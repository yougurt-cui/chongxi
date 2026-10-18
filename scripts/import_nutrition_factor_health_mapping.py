#!/usr/bin/env python3
"""Import the six-factor health-focus workbook into the local feature database."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import pymysql
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app_config import get_feature_mysql_config  # noqa: E402


TABLE = "catfood_nutrition_factor_health_mapping"
SHEET = "nutrition_factor_health_mapping"
VERSION = "v1"
EXPECTED_HEADERS = (
    "factor_code", "factor_name", "factor_status", "health_focus",
    "health_mechanism", "user_observation",
)
FACTOR_CODES = {
    "PROTEIN_PRESSURE", "PROTEIN_QUALITY", "FAT_BURDEN",
    "FIBER_SUPPORT", "CARB_BURDEN", "PREBIOTIC_SUPPORT",
}
FACTOR_STATUSES = {"high", "low"}


def utc_now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def clean(value, limit: int | None = None) -> str:
    result = str(value or "").strip()
    return result[:limit] if limit else result


def mapping_code(row: dict) -> str:
    identity = "|".join(
        (row["factor_code"], row["factor_status"], row["health_focus"])
    )
    return "NFH_" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:24].upper()


def read_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise ValueError(f"缺少工作表: {SHEET}")
    worksheet = workbook[SHEET]
    values = worksheet.iter_rows(values_only=True)
    headers = tuple(clean(value) for value in next(values, ()))
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"表头不匹配: expected={EXPECTED_HEADERS} actual={headers}")

    rows = []
    seen = set()
    for source_row_number, values_row in enumerate(values, 2):
        if not any(value not in (None, "") for value in values_row):
            continue
        row = dict(zip(headers, values_row))
        row = {
            "factor_code": clean(row["factor_code"], 64).upper(),
            "factor_name": clean(row["factor_name"], 64),
            "factor_status": clean(row["factor_status"], 32).lower(),
            "health_focus": clean(row["health_focus"], 128),
            "health_mechanism": clean(row["health_mechanism"], 4000),
            "user_observation": clean(row["user_observation"], 2000),
            "source_row_number": source_row_number,
        }
        if row["factor_code"] not in FACTOR_CODES:
            raise ValueError(f"第 {source_row_number} 行 factor_code 无效: {row['factor_code']}")
        if row["factor_status"] not in FACTOR_STATUSES:
            raise ValueError(f"第 {source_row_number} 行 factor_status 无效: {row['factor_status']}")
        for required in ("factor_name", "health_focus", "health_mechanism", "user_observation"):
            if not row[required]:
                raise ValueError(f"第 {source_row_number} 行 {required} 不能为空")
        identity = (row["factor_code"], row["factor_status"], row["health_focus"])
        if identity in seen:
            raise ValueError(f"第 {source_row_number} 行存在重复映射: {identity}")
        seen.add(identity)
        row["mapping_code"] = mapping_code(row)
        rows.append(row)
    if not rows:
        raise ValueError("工作表没有可导入数据")
    return rows


def ensure_table(cursor) -> None:
    cursor.execute(
        f"""
        CREATE TABLE IF NOT EXISTS `{TABLE}` (
          id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
          mapping_code VARCHAR(32) NOT NULL,
          factor_code VARCHAR(64) NOT NULL,
          factor_name VARCHAR(64) NOT NULL,
          factor_status VARCHAR(32) NOT NULL,
          health_focus VARCHAR(128) NOT NULL,
          health_mechanism TEXT NOT NULL,
          user_observation TEXT NOT NULL,
          relation_type VARCHAR(32) NULL,
          knowledge_version VARCHAR(32) NOT NULL,
          source_batch_id VARCHAR(64) NOT NULL,
          source_sheet_name VARCHAR(255) NOT NULL,
          source_row_number INT NOT NULL,
          review_status VARCHAR(32) NOT NULL DEFAULT 'pending',
          status VARCHAR(16) NOT NULL DEFAULT 'active',
          created_at DATETIME NOT NULL,
          updated_at DATETIME NOT NULL,
          PRIMARY KEY (id),
          UNIQUE KEY uk_factor_health_version
            (factor_code, factor_status, health_focus, knowledge_version),
          UNIQUE KEY uk_factor_health_mapping_code (mapping_code, knowledge_version),
          KEY idx_factor_health_lookup (factor_code, factor_status, status),
          KEY idx_factor_health_focus (health_focus, status),
          KEY idx_factor_health_review (review_status, status)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
        """
    )


def import_workbook(path: Path, *, version: str = VERSION) -> dict:
    rows = read_rows(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    batch_id = f"nutrition-factor-{version}-{digest[:16]}"
    now = utc_now()
    config = get_feature_mysql_config()
    with pymysql.connect(
        **config, cursorclass=pymysql.cursors.DictCursor
    ) as connection:
        with connection.cursor() as cursor:
            ensure_table(cursor)
            cursor.execute(
                """
                INSERT INTO nutrition_kb_import_batch (
                  batch_id,source_filename,source_sha256,target_knowledge_version,
                  import_status,total_rows,valid_rows,invalid_rows,imported_rows,
                  created_by,created_at,completed_at
                ) VALUES (%s,%s,%s,%s,'imported',%s,%s,0,%s,'local:openpyxl',%s,%s)
                ON DUPLICATE KEY UPDATE
                  source_filename=VALUES(source_filename),source_sha256=VALUES(source_sha256),
                  target_knowledge_version=VALUES(target_knowledge_version),
                  import_status='imported',total_rows=VALUES(total_rows),
                  valid_rows=VALUES(valid_rows),invalid_rows=0,
                  imported_rows=VALUES(imported_rows),completed_at=VALUES(completed_at)
                """,
                (batch_id, path.name, digest, version, len(rows), len(rows), len(rows), now, now),
            )
            for row in rows:
                raw = {key: row[key] for key in EXPECTED_HEADERS}
                normalized = {**raw, "mapping_code": row["mapping_code"]}
                cursor.execute(
                    """
                    INSERT INTO nutrition_kb_staging_row (
                      batch_id,sheet_name,source_row_number,entity_type,entity_key,
                      raw_row_json,normalized_row_json,validation_status,
                      validation_errors_json,target_table,created_at,updated_at
                    ) VALUES (%s,%s,%s,'factor_health_mapping',%s,%s,%s,'valid',JSON_ARRAY(),%s,%s,%s)
                    ON DUPLICATE KEY UPDATE
                      entity_type=VALUES(entity_type),entity_key=VALUES(entity_key),
                      raw_row_json=VALUES(raw_row_json),normalized_row_json=VALUES(normalized_row_json),
                      validation_status='valid',validation_errors_json=JSON_ARRAY(),
                      target_table=VALUES(target_table),updated_at=VALUES(updated_at)
                    """,
                    (
                        batch_id, SHEET, row["source_row_number"], row["mapping_code"],
                        json.dumps(raw, ensure_ascii=False),
                        json.dumps(normalized, ensure_ascii=False), TABLE, now, now,
                    ),
                )
                cursor.execute(
                    f"""
                    INSERT INTO `{TABLE}` (
                      mapping_code,factor_code,factor_name,factor_status,health_focus,
                      health_mechanism,user_observation,relation_type,knowledge_version,
                      source_batch_id,source_sheet_name,source_row_number,
                      review_status,status,created_at,updated_at
                    ) VALUES (%s,%s,%s,%s,%s,%s,%s,NULL,%s,%s,%s,%s,'pending','active',%s,%s)
                    ON DUPLICATE KEY UPDATE
                      mapping_code=VALUES(mapping_code),factor_name=VALUES(factor_name),
                      health_mechanism=VALUES(health_mechanism),
                      user_observation=VALUES(user_observation),
                      source_batch_id=VALUES(source_batch_id),
                      source_sheet_name=VALUES(source_sheet_name),
                      source_row_number=VALUES(source_row_number),
                      review_status='pending',status='active',updated_at=VALUES(updated_at)
                    """,
                    (
                        row["mapping_code"], row["factor_code"], row["factor_name"],
                        row["factor_status"], row["health_focus"], row["health_mechanism"],
                        row["user_observation"], version, batch_id, SHEET,
                        row["source_row_number"], now, now,
                    ),
                )
            cursor.execute(
                f"SELECT COUNT(*) AS count FROM `{TABLE}` WHERE knowledge_version=%s",
                (version,),
            )
            imported_count = int(cursor.fetchone()["count"])
        connection.commit()
    return {
        "ok": True,
        "database": config.get("database"),
        "table": TABLE,
        "knowledge_version": version,
        "batch_id": batch_id,
        "source_rows": len(rows),
        "imported_count": imported_count,
        "review_status": "pending",
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--version", default=VERSION)
    args = parser.parse_args()
    if not args.workbook.is_file():
        raise SystemExit(f"文件不存在: {args.workbook}")
    print(json.dumps(import_workbook(args.workbook, version=args.version), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
