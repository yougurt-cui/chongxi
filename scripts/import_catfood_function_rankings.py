#!/usr/bin/env python3
"""Import JD functional cat-food ranking candidates from the supplied workbook."""

from __future__ import annotations

import argparse
import hashlib
import sys
from pathlib import Path

import pymysql
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from app_config import get_mysql_config


CATEGORIES = ("肠胃软便", "美毛", "泌尿", "化毛", "低敏", "控重")
HEADERS = (
    "序号", "品牌", "SKU/商品名", "评论/热度量级", "功效证据", "处方粮",
    "排名性质", "平台", "来源URL", "抓取日期", "备注",
)


def _text(value: object) -> str:
    return str(value or "").strip()


def read_rows(workbook_path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    rows: list[dict[str, object]] = []
    try:
        for category in CATEGORIES:
            if category not in workbook.sheetnames:
                raise ValueError(f"缺少工作表：{category}")
            sheet = workbook[category]
            headers = tuple(_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
            if headers[: len(HEADERS)] != HEADERS:
                raise ValueError(f"工作表 {category} 字段与预期不一致")
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if not any(value is not None and _text(value) for value in values):
                    continue
                row = dict(zip(HEADERS, values[: len(HEADERS)]))
                rank = int(row["序号"])
                key_text = "|".join((category, str(rank), _text(row["品牌"]), _text(row["SKU/商品名"]), _text(row["抓取日期"])))
                rows.append({
                    "source_row_key": hashlib.sha256(key_text.encode("utf-8")).hexdigest(),
                    "effect_category": category,
                    "rank_in_list": rank,
                    "brand": _text(row["品牌"]),
                    "product_name": _text(row["SKU/商品名"]),
                    "heat_level": _text(row["评论/热度量级"]) or None,
                    "effect_evidence": _text(row["功效证据"]),
                    "is_prescription": 1 if _text(row["处方粮"]) == "是" else 0,
                    "ranking_nature": _text(row["排名性质"]),
                    "platform": _text(row["平台"]),
                    "source_url": _text(row["来源URL"]),
                    "captured_date": row["抓取日期"],
                    "notes": _text(row["备注"]) or None,
                })
    finally:
        workbook.close()
    return rows


def import_rows(workbook_path: Path) -> int:
    rows = read_rows(workbook_path)
    if len(rows) != 91:
        raise ValueError(f"预期导入 91 条，实际读取 {len(rows)} 条")
    connection = pymysql.connect(
        **get_mysql_config(database="csv_labeling"),
        cursorclass=pymysql.cursors.DictCursor,
        autocommit=False,
    )
    try:
        with connection.cursor() as cursor:
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS catfood_function_sku_ranking (
                    id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
                    source_row_key CHAR(64) NOT NULL,
                    effect_category VARCHAR(64) NOT NULL,
                    rank_in_list INT UNSIGNED NOT NULL,
                    brand VARCHAR(255) NOT NULL,
                    product_name VARCHAR(512) NOT NULL,
                    heat_level VARCHAR(64) NULL,
                    effect_evidence VARCHAR(512) NOT NULL,
                    is_prescription TINYINT(1) NOT NULL DEFAULT 0,
                    ranking_nature VARCHAR(128) NOT NULL,
                    platform VARCHAR(64) NOT NULL,
                    source_url TEXT NOT NULL,
                    captured_date DATE NOT NULL,
                    notes VARCHAR(512) NULL,
                    source_file VARCHAR(255) NOT NULL,
                    imported_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
                    PRIMARY KEY (id),
                    UNIQUE KEY uq_function_sku_source_row (source_row_key),
                    KEY idx_function_sku_category_rank (effect_category, rank_in_list),
                    KEY idx_function_sku_brand (brand),
                    KEY idx_function_sku_date (captured_date)
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci
            """)
            values = [tuple(row[key] for key in (
                "source_row_key", "effect_category", "rank_in_list", "brand", "product_name",
                "heat_level", "effect_evidence", "is_prescription", "ranking_nature", "platform",
                "source_url", "captured_date", "notes",
            )) + (workbook_path.name,) for row in rows]
            cursor.executemany("""
                INSERT INTO catfood_function_sku_ranking (
                    source_row_key, effect_category, rank_in_list, brand, product_name,
                    heat_level, effect_evidence, is_prescription, ranking_nature, platform,
                    source_url, captured_date, notes, source_file
                ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                ON DUPLICATE KEY UPDATE
                    effect_category=VALUES(effect_category), rank_in_list=VALUES(rank_in_list),
                    brand=VALUES(brand), product_name=VALUES(product_name), heat_level=VALUES(heat_level),
                    effect_evidence=VALUES(effect_evidence), is_prescription=VALUES(is_prescription),
                    ranking_nature=VALUES(ranking_nature), platform=VALUES(platform),
                    source_url=VALUES(source_url), captured_date=VALUES(captured_date),
                    notes=VALUES(notes), source_file=VALUES(source_file)
            """, values)
        connection.commit()
        return len(values)
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    count = import_rows(args.workbook.resolve())
    print(f"Imported {count} functional SKU ranking rows")


if __name__ == "__main__":
    main()
